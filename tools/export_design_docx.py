"""Export the main Markdown design to an editable Word document.

Use the managed document runtime. No project or system package installation is needed.
This export does not verify layout; render and review the resulting document separately.
"""

from pathlib import Path
import math
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "智能实验室终端设计开发文档.md"
QA = ROOT / ".qa"
OUTPUT = QA / SOURCE.with_suffix(".docx").name


def plain(text):
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    return text.replace("**", "").replace("`", "")


def font_style(style, size, bold=False, east="宋体", latin="Calibri"):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), east)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)


def pil_font(size):
    return ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", size)


def diagram(kind):
    path = QA / f"{kind}.png"
    if kind == "workflow":
        labels = ["登录与选实验", "首拍与质量检查", "冻结仪器基线", "资料与实验", "结束与收台向导", "终拍与系统分析", "保存并同步", "教师查看与导出"]
        im = Image.new("RGB", (2200, 670), "white")
        d = ImageDraw.Draw(im)
        positions = [(60 + i * 540, 80) for i in range(4)] + [(60 + i * 540, 390) for i in range(4)]
        for i, (x, y) in enumerate(positions):
            d.rounded_rectangle((x, y, x + 450, y + 160), 16, fill="#F3F4F5", outline="#808080", width=3)
            label = f"{i + 1}  {labels[i]}"
            d.text((x + 225, y + 80), label, font=pil_font(36), fill="black", anchor="mm")
            if i not in (3, 7):
                arrow(d, (x + 455, y + 80), (x + 530, y + 80))
        d.line([(1960, 248), (1960, 300), (285, 300), (285, 360)], fill="#555555", width=4)
        arrow(d, (285, 360), (285, 382))
    else:
        im = Image.new("RGB", (2200, 1000), "white")
        d = ImageDraw.Draw(im)
        nodes = {
            "camera": (80, 100, "固定摄像头"), "ui": (80, 365, "触摸屏 Web UI"),
            "agent": (670, 215, "终端 Agent"), "vision": (670, 600, "检测与 AnchorDiff"),
            "local": (80, 700, "SQLite 与证据"), "api": (1270, 215, "HTTPS 同步 API"),
            "store": (1800, 100, "数据库与存储"), "teacher": (1270, 600, "教师只读后台"),
            "export": (1800, 600, "报告导出"),
        }
        for x, y, label in nodes.values():
            d.rounded_rectangle((x, y, x + 330, y + 145), 16, fill="#F3F4F5", outline="#808080", width=3)
            size = 32 if len(label) < 13 else 28
            d.text((x + 165, y + 72), label, font=pil_font(size), fill="black", anchor="mm")
        for a, b in [("camera", "agent"), ("ui", "agent"), ("agent", "api"), ("api", "store"), ("teacher", "api"), ("api", "export"), ("agent", "vision"), ("agent", "local")]:
            ax, ay, _ = nodes[a]
            bx, by, _ = nodes[b]
            if bx > ax:
                arrow(d, (ax + 338, ay + 72), (bx - 10, by + 72))
            elif by > ay:
                arrow(d, (ax + 165, ay + 150), (bx + 165, by - 12))
            else:
                arrow(d, (ax + 165, ay - 8), (bx + 165, by + 154))
        d.text((60, 935), "端侧完成分析与离线保存，后台校验、归档、查看和导出", font=pil_font(32), fill="black")
    im.save(path)
    return path


def arrow(draw, start, end):
    draw.line([start, end], fill="#555555", width=4)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 20
    points = [end] + [(end[0] - length * math.cos(angle + a), end[1] - length * math.sin(angle + a)) for a in (-0.4, 0.4)]
    draw.polygon(points, fill="#555555")


def table(doc, rows):
    if not rows:
        return
    cols = len(rows[0])
    t = doc.add_table(rows=1, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if cols == 2:
        proportions = [0.26, 0.74]
    elif cols == 3:
        proportions = [0.20, 0.34, 0.46]
        if rows[0][0] in ("编号", "阶段"):
            proportions = [0.13, 0.38, 0.49]
    elif cols == 4:
        proportions = [0.17, 0.26, 0.30, 0.27]
    else:
        proportions = [1 / cols] * cols
    for col, weight in zip(t.columns, proportions):
        col.width = Cm(17 * weight)
    for ridx, row in enumerate(rows):
        cells = t.rows[0].cells if ridx == 0 else t.add_row().cells
        for idx, cell in enumerate(cells):
            cell.width = Cm(17 * proportions[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.12
            p.paragraph_format.keep_with_next = False
            run = p.add_run(plain(row[idx] if idx < len(row) else ""))
            run.font.size = Pt(9.5)
            run.font.bold = ridx == 0
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                item = OxmlElement(f"w:{side}")
                item.set(qn("w:w"), "85")
                item.set(qn("w:type"), "dxa")
                margins.append(item)
            props.append(margins)
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                item = OxmlElement(f"w:{side}")
                item.set(qn("w:val"), "single")
                item.set(qn("w:sz"), "4")
                item.set(qn("w:color"), "CCCCCC")
                borders.append(item)
            props.append(borders)
            if ridx == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "F2F2F2")
                props.append(shade)
        trpr = t.rows[ridx]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
    header = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(header)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def main():
    QA.mkdir(exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.9)
    sec.left_margin, sec.right_margin = Cm(2), Cm(2)
    font_style(doc.styles["Normal"], 10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.22
    font_style(doc.styles["Title"], 24, True, east="微软雅黑")
    for level, size in [(1, 16), (2, 12), (3, 11)]:
        style = doc.styles[f"Heading {level}"]
        font_style(style, size, True, east="微软雅黑")
        style.paragraph_format.space_before = Pt(13 if level == 1 else 9)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    if "Code" not in doc.styles:
        doc.styles.add_style("Code", 1)
    font_style(doc.styles["Code"], 8.5, latin="Consolas")
    doc.styles["Code"].paragraph_format.space_after = Pt(1)
    doc.styles["Code"].paragraph_format.line_spacing = 1.05
    doc.core_properties.author = "Smart Lab Terminal"
    doc.core_properties.last_modified_by = "Smart Lab Terminal"
    doc.core_properties.title = "智能实验室终端设计开发文档"
    doc.core_properties.subject = "设计基线 0.1.0"
    doc.core_properties.keywords = "终端 设计 目标检测 AnchorDiff RK3588"
    doc.core_properties.comments = ""
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    idx = 0
    while idx < len(lines):
        line = lines[idx]
        if not line.strip():
            idx += 1
            continue
        if line.startswith("```"):
            language = line[3:].strip()
            block = []
            idx += 1
            while idx < len(lines) and not lines[idx].startswith("```"):
                block.append(lines[idx])
                idx += 1
            if language == "mermaid":
                kind = "workflow" if any("登录" in value for value in block) else "architecture"
                p = doc.add_paragraph()
                p.add_run().add_picture(str(diagram(kind)), width=Cm(17))
                p.paragraph_format.keep_with_next = True
                p = doc.add_paragraph("图 1 业务流程" if kind == "workflow" else "图 2 总体架构")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                for value in block:
                    doc.add_paragraph(value, style="Code")
            idx += 1
            continue
        if line.startswith("|"):
            rows = []
            while idx < len(lines) and lines[idx].startswith("|"):
                value = lines[idx]
                cells = [cell.strip() for cell in value.strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                    rows.append(cells)
                idx += 1
            table(doc, rows)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            level = len(heading[1])
            label = re.sub(r"(?<=\d)\.(?=\d)", " ", plain(heading[2]))
            doc.add_paragraph(label, style="Title" if level == 1 else f"Heading {min(level - 1, 3)}")
        else:
            doc.add_paragraph(plain(line))
        idx += 1
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("智能实验室终端  0.1.0  |  ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.save(OUTPUT)
    check = Document(OUTPUT)
    print(f"Saved {OUTPUT.name}; paragraphs={len(check.paragraphs)}, tables={len(check.tables)}")


if __name__ == "__main__":
    main()
